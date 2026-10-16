"""Servidor SIGACEcyTEO: API, panel encargado, PostgreSQL/SQLite de desarrollo."""
from __future__ import annotations
import asyncio, base64, csv, hashlib, hmac, io, os, secrets, time
from datetime import datetime, timedelta, timezone
from typing import Optional
import qrcode, uvicorn
from dotenv import load_dotenv
from fastapi import FastAPI, Depends, Header, HTTPException, Request, WebSocket, WebSocketDisconnect
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, DateTime, Float, ForeignKey, Integer, String, Text, create_engine, select
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, sessionmaker
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

load_dotenv()
DATABASE_URL=os.getenv("DATABASE_URL","sqlite:///./sigacecyteo_dev.db")
HOST=os.getenv("SIGA_HOST","0.0.0.0"); PORT=int(os.getenv("SIGA_PORT","8080"))
ITERATIONS=310000; TOKEN_HOURS=int(os.getenv("SIGA_TOKEN_HOURS","8"))
engine=create_engine(DATABASE_URL,connect_args={"check_same_thread":False} if DATABASE_URL.startswith("sqlite") else {},pool_pre_ping=True)
SessionLocal=sessionmaker(bind=engine,expire_on_commit=False)
class Base(DeclarativeBase): pass
def now(): return datetime.now(timezone.utc)
def ts(x): return x.isoformat() if x else ""
def hashpw(password):
    salt=secrets.token_bytes(16); digest=hashlib.pbkdf2_hmac("sha256",password.encode(),salt,ITERATIONS)
    return "pbkdf2_sha256$"+str(ITERATIONS)+"$"+salt.hex()+"$"+digest.hex()
def checkpw(password,stored):
    try:
        alg,n,salt,digest=stored.split("$")
        calc=hashlib.pbkdf2_hmac("sha256",password.encode(),bytes.fromhex(salt),int(n)).hex()
        return alg=="pbkdf2_sha256" and hmac.compare_digest(calc,digest)
    except Exception: return False

class Admin(Base):
    __tablename__="admin_users"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    username:Mapped[str]=mapped_column(String(80),unique=True,index=True)
    password_hash:Mapped[str]=mapped_column(String(256))
    role:Mapped[str]=mapped_column(String(30),default="Administrador")
    active:Mapped[bool]=mapped_column(Boolean,default=True)
class Student(Base):
    __tablename__="students"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    matricula:Mapped[str]=mapped_column(String(40),unique=True,index=True)
    nombre:Mapped[str]=mapped_column(String(160))
    password_hash:Mapped[str]=mapped_column(String(256))
    active:Mapped[bool]=mapped_column(Boolean,default=True)
class Computer(Base):
    __tablename__="computers"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    code:Mapped[str]=mapped_column(String(60),unique=True,index=True)
    hostname:Mapped[str]=mapped_column(String(160),default="")
    ip_address:Mapped[str]=mapped_column(String(80),default="")
    state:Mapped[str]=mapped_column(String(30),default="Libre")
    last_seen:Mapped[Optional[datetime]]=mapped_column(DateTime(timezone=True),nullable=True)
    os_name:Mapped[str]=mapped_column(String(120),default="")
    cpu:Mapped[Optional[float]]=mapped_column(Float,nullable=True)
    memory:Mapped[Optional[float]]=mapped_column(Float,nullable=True)
    disk:Mapped[Optional[float]]=mapped_column(Float,nullable=True)
class LabSession(Base):
    __tablename__="lab_sessions"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    offline_id:Mapped[Optional[str]]=mapped_column(String(80),unique=True,nullable=True)
    student_id:Mapped[int]=mapped_column(ForeignKey("students.id"))
    computer_id:Mapped[int]=mapped_column(ForeignKey("computers.id"))
    started_at:Mapped[datetime]=mapped_column(DateTime(timezone=True),default=now)
    ended_at:Mapped[Optional[datetime]]=mapped_column(DateTime(timezone=True),nullable=True)
    state:Mapped[str]=mapped_column(String(40),default="Activa")
    mode:Mapped[str]=mapped_column(String(20),default="password")
class Incident(Base):
    __tablename__="incidents"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    student_id:Mapped[int]=mapped_column(ForeignKey("students.id"))
    computer_id:Mapped[int]=mapped_column(ForeignKey("computers.id"))
    session_id:Mapped[Optional[int]]=mapped_column(ForeignKey("lab_sessions.id"),nullable=True)
    event_id:Mapped[Optional[str]]=mapped_column(String(80),unique=True,nullable=True)
    category:Mapped[str]=mapped_column(String(80))
    description:Mapped[str]=mapped_column(Text)
    state:Mapped[str]=mapped_column(String(30),default="Abierta")
    created_at:Mapped[datetime]=mapped_column(DateTime(timezone=True),default=now)
    resolved_at:Mapped[Optional[datetime]]=mapped_column(DateTime(timezone=True),nullable=True)
    resolution_note:Mapped[str]=mapped_column(Text,default="")
class Command(Base):
    __tablename__="remote_commands"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    computer_id:Mapped[int]=mapped_column(ForeignKey("computers.id"))
    command:Mapped[str]=mapped_column(String(30))
    status:Mapped[str]=mapped_column(String(30),default="Pendiente")
    created_at:Mapped[datetime]=mapped_column(DateTime(timezone=True),default=now)
    sent_at:Mapped[Optional[datetime]]=mapped_column(DateTime(timezone=True),nullable=True)
    executed_at:Mapped[Optional[datetime]]=mapped_column(DateTime(timezone=True),nullable=True)
class QRToken(Base):
    __tablename__="qr_tokens"
    id:Mapped[int]=mapped_column(Integer,primary_key=True)
    token_hash:Mapped[str]=mapped_column(String(64),unique=True,index=True)
    student_id:Mapped[int]=mapped_column(ForeignKey("students.id"))
    expires_at:Mapped[datetime]=mapped_column(DateTime(timezone=True))
    consumed_at:Mapped[Optional[datetime]]=mapped_column(DateTime(timezone=True),nullable=True)

TOKENS:dict[str,dict]={}
app=FastAPI(title="SIGACEcyTEO Server",version="1.0.0",docs_url="/api/docs",redoc_url=None)
def getdb():
    db=SessionLocal()
    try: yield db
    finally: db.close()
@app.on_event("startup")
def startup():
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        # Tokens live in memory and expire after a process restart; close stale DB sessions to
        # avoid leaving computers permanently "occupied" after a planned server reboot.
        stale=db.scalars(select(LabSession).where(LabSession.ended_at.is_(None))).all()
        if stale:
            affected=set()
            for session in stale:
                session.ended_at=now()
                session.state="Cerrada por reinicio del servidor"
                affected.add(session.computer_id)
            for computer_id in affected:
                computer=db.get(Computer,computer_id)
                if computer and computer.state=="En uso":
                    computer.state="Libre"
            db.commit()
        username=os.getenv("SIGA_ADMIN_USER","").strip(); password=os.getenv("SIGA_ADMIN_PASSWORD","")
        if username and password and not db.scalar(select(Admin).where(Admin.username==username)):
            db.add(Admin(username=username,password_hash=hashpw(password),role="Administrador",active=True))
        if os.getenv("SIGA_SEED_DEMO","false").lower() in ("1","true","yes"):
            for mat,name,pw in (("2026001","Alumno Demo 1","cecyte123"),("2026002","Alumno Demo 2","cecyte123")):
                if not db.scalar(select(Student).where(Student.matricula==mat)):
                    db.add(Student(matricula=mat,nombre=name,password_hash=hashpw(pw)))
        db.commit()
    if not os.getenv("SIGA_ADMIN_USER") or not os.getenv("SIGA_ADMIN_PASSWORD"):
        print("AVISO: configura SIGA_ADMIN_USER y SIGA_ADMIN_PASSWORD en .env para inicializar el administrador.")
def issue(kind,ident,more=None):
    t=secrets.token_urlsafe(40); TOKENS[t]={"kind":kind,"id":ident,"expires":time.time()+TOKEN_HOURS*3600,**(more or {})}; return t
def auth(authorization,kind):
    if not authorization or not authorization.startswith("Bearer "): raise HTTPException(401,"Inicia sesión.")
    t=authorization[7:].strip(); d=TOKENS.get(t)
    if not d or d.get("expires",0)<time.time() or d.get("kind")!=kind:
        TOKENS.pop(t,None); raise HTTPException(401,"Sesión vencida; vuelve a iniciar sesión.")
    return d
def admin_auth(authorization:str|None=Header(default=None)): return auth(authorization,"admin")
def client_auth(authorization:str|None=Header(default=None),db=Depends(getdb)):
    d=auth(authorization,"client"); s=db.get(LabSession,d["session_id"])
    if not s or s.ended_at: raise HTTPException(401,"La sesión del equipo ya terminó.")
    return d

class AdminLogin(BaseModel): username:str; password:str
class StudentLogin(BaseModel): matricula:str; password:str; computer_code:str=Field(min_length=1,max_length=60)
class QRLogin(BaseModel): qr_token:str; computer_code:str=Field(min_length=1,max_length=60)
class StudentNew(BaseModel): matricula:str=Field(min_length=2,max_length=40); nombre:str=Field(min_length=2,max_length=160); password:str=Field(min_length=10,max_length=256)
class QRNew(BaseModel): matricula:str
class IncidentNew(BaseModel): category:str=Field(min_length=2,max_length=80); description:str=Field(min_length=3,max_length=2000)
class CommandNew(BaseModel): computer_code:str; command:str
class ResolveBody(BaseModel): note:str=""
class Telemetry(BaseModel): hostname:str=""; os_name:str=""; cpu:float|None=None; memory:float|None=None; disk:float|None=None
class OfflineSync(BaseModel): events:list[dict]=Field(max_length=500)

def computer_get(db,code):
    code=code.strip().upper()
    if not code: raise HTTPException(400,"Falta código de equipo.")
    c=db.scalar(select(Computer).where(Computer.code==code))
    if not c: c=Computer(code=code,state="Libre"); db.add(c); db.flush()
    return c
def new_session(db,student,code,mode,ip):
    c=computer_get(db,code)
    if c.state=="Bloqueado": raise HTTPException(409,"El equipo está bloqueado por el encargado.")
    if db.scalar(select(LabSession).where(LabSession.computer_id==c.id,LabSession.ended_at.is_(None))): raise HTTPException(409,"Este equipo ya tiene sesión activa.")
    if db.scalar(select(LabSession).where(LabSession.student_id==student.id,LabSession.ended_at.is_(None))): raise HTTPException(409,"La matrícula ya tiene sesión activa.")
    s=LabSession(student_id=student.id,computer_id=c.id,started_at=now(),state="Activa",mode=mode); db.add(s)
    c.state="En uso"; c.last_seen=now(); c.ip_address=ip; db.commit(); db.refresh(s)
    token=issue("client",student.id,{"session_id":s.id,"computer_id":c.id,"computer_code":c.code})
    return {"access_token":token,"token_type":"bearer","session_id":s.id,"matricula":student.matricula,"nombre":student.nombre,"computer_code":c.code,"started_at":ts(s.started_at)}

@app.get("/api/health")
def health(): return {"ok":True,"service":"SIGACEcyTEO","time":ts(now())}
@app.post("/api/admin/login")
def admin_login(body:AdminLogin,db=Depends(getdb)):
    u=db.scalar(select(Admin).where(Admin.username==body.username.strip(),Admin.active.is_(True)))
    if not u or not checkpw(body.password,u.password_hash): raise HTTPException(401,"Usuario o contraseña incorrectos.")
    return {"access_token":issue("admin",u.id,{"username":u.username,"role":u.role}),"username":u.username,"role":u.role}
@app.post("/api/auth/login")
def login(body:StudentLogin,request:Request,db=Depends(getdb)):
    s=db.scalar(select(Student).where(Student.matricula==body.matricula.strip(),Student.active.is_(True)))
    if not s or not checkpw(body.password,s.password_hash): raise HTTPException(401,"Matrícula o contraseña incorrectas.")
    return new_session(db,s,body.computer_code,"password",request.client.host if request.client else "")
@app.post("/api/auth/qr-login")
def qrlogin(body:QRLogin,request:Request,db=Depends(getdb)):
    raw=body.qr_token.removeprefix("SIGACEcyTEO:")
    digest=hashlib.sha256(raw.encode()).hexdigest()
    q=db.scalar(select(QRToken).where(QRToken.token_hash==digest,QRToken.consumed_at.is_(None)))
    if not q or q.expires_at.replace(tzinfo=timezone.utc)<now(): raise HTTPException(401,"QR vencido o utilizado; pide otro al encargado.")
    student=db.get(Student,q.student_id)
    if not student or not student.active: raise HTTPException(401,"Alumno inactivo.")
    q.consumed_at=now(); result=new_session(db,student,body.computer_code,"qr",request.client.host if request.client else ""); db.commit(); return result
@app.post("/api/client/heartbeat")
def heartbeat(body:Telemetry,request:Request,who=Depends(client_auth),db=Depends(getdb)):
    c=db.get(Computer,who["computer_id"])
    c.hostname=body.hostname[:160]; c.os_name=body.os_name[:120]; c.cpu=body.cpu; c.memory=body.memory; c.disk=body.disk
    c.ip_address=request.client.host if request.client else c.ip_address; c.last_seen=now(); db.commit()
    return {"ok":True,"state":c.state,"server_time":ts(now())}
@app.post("/api/client/incidents")
def report_incident(body:IncidentNew,who=Depends(client_auth),db=Depends(getdb)):
    s=db.get(LabSession,who["session_id"])
    if not s or s.ended_at: raise HTTPException(401,"Sesión finalizada.")
    x=Incident(student_id=s.student_id,computer_id=s.computer_id,session_id=s.id,category=body.category.strip(),description=body.description.strip())
    db.add(x); db.commit(); db.refresh(x); return {"ok":True,"incident_id":x.id,"state":x.state}
@app.post("/api/client/logout")
def client_logout(who=Depends(client_auth),authorization:str|None=Header(default=None),db=Depends(getdb)):
    s=db.get(LabSession,who["session_id"])
    if s and not s.ended_at:
        s.ended_at=now(); s.state="Cerrada"; c=db.get(Computer,s.computer_id)
        if c: c.state="Libre" if c.state!="Bloqueado" else c.state; c.last_seen=now()
        db.commit()
    if authorization: TOKENS.pop(authorization[7:].strip(),None)
    return {"ok":True}
@app.post("/api/client/command-ack")
def command_ack(body:dict,who=Depends(client_auth),db=Depends(getdb)):
    c=db.get(Command,int(body.get("command_id",0)))
    if not c or c.computer_id!=who["computer_id"]: raise HTTPException(404,"Comando no encontrado.")
    c.status="Ejecutado"; c.executed_at=now(); db.commit(); return {"ok":True}
@app.post("/api/client/sync-offline")
def sync_offline(body:OfflineSync,who=Depends(client_auth),db=Depends(getdb)):
    live=db.get(LabSession,who["session_id"]); count=0
    for ev in body.events:
        try:
            kind=ev.get("type"); offline_id=str(ev.get("offline_id",""))[:80]; at=datetime.fromisoformat(ev.get("at",""))
            if kind=="session_start":
                if offline_id and not db.scalar(select(LabSession).where(LabSession.offline_id==offline_id)):
                    c=computer_get(db,ev.get("computer_code",who["computer_code"]))
                    db.add(LabSession(offline_id=offline_id,student_id=who["id"],computer_id=c.id,started_at=at,state="Modo offline sincronizado",mode="offline")); count+=1
            elif kind=="incident":
                event_id=str(ev.get("event_id",""))[:80]
                if event_id and db.scalar(select(Incident).where(Incident.event_id==event_id)): continue
                code=str(ev.get("computer_code",who["computer_code"])).upper()
                c=db.scalar(select(Computer).where(Computer.code==code)) or db.get(Computer,who["computer_id"])
                old=db.scalar(select(LabSession).where(LabSession.offline_id==offline_id,LabSession.student_id==who["id"])) if offline_id else None
                db.add(Incident(student_id=who["id"],computer_id=c.id,session_id=old.id if old else live.id,event_id=event_id or None,category=str(ev.get("category","Incidencia offline"))[:80],description=str(ev.get("description",""))[:2000],created_at=at)); count+=1
            elif kind=="session_end":
                old=db.scalar(select(LabSession).where(LabSession.offline_id==offline_id,LabSession.student_id==who["id"])) if offline_id else None
                if old and not old.ended_at: old.ended_at=at; old.state="Cerrada (sin conexión)"; count+=1
        except (ValueError,TypeError,KeyError): continue
    db.commit(); return {"ok":True,"synced":count}
@app.get("/api/admin/state")
def admin_state(who=Depends(admin_auth),db=Depends(getdb)):
    computers=db.scalars(select(Computer).order_by(Computer.code)).all(); students=db.scalars(select(Student).order_by(Student.matricula)).all()
    sessions=db.scalars(select(LabSession).order_by(LabSession.id.desc()).limit(500)).all(); incidents=db.scalars(select(Incident).order_by(Incident.id.desc()).limit(500)).all(); commands=db.scalars(select(Command).order_by(Command.id.desc()).limit(100)).all()
    sm={s.id:s for s in students}; cm={c.id:c for c in computers}
    return {"computers":[{"id":c.id,"code":c.code,"hostname":c.hostname,"ip_address":c.ip_address,"state":c.state,"last_seen":ts(c.last_seen),"os_name":c.os_name,"cpu":c.cpu,"memory":c.memory,"disk":c.disk} for c in computers],
    "students":[{"id":s.id,"matricula":s.matricula,"nombre":s.nombre,"active":s.active} for s in students],
    "sessions":[{"id":s.id,"matricula":sm[s.student_id].matricula if s.student_id in sm else "","nombre":sm[s.student_id].nombre if s.student_id in sm else "","computer_code":cm[s.computer_id].code if s.computer_id in cm else "","started_at":ts(s.started_at),"ended_at":ts(s.ended_at),"state":s.state,"mode":s.mode} for s in sessions],
    "incidents":[{"id":x.id,"matricula":sm[x.student_id].matricula if x.student_id in sm else "","computer_code":cm[x.computer_id].code if x.computer_id in cm else "","category":x.category,"description":x.description,"state":x.state,"created_at":ts(x.created_at),"resolved_at":ts(x.resolved_at),"resolution_note":x.resolution_note} for x in incidents],
    "commands":[{"id":x.id,"computer_code":cm[x.computer_id].code if x.computer_id in cm else "","command":x.command,"status":x.status,"created_at":ts(x.created_at),"executed_at":ts(x.executed_at)} for x in commands]}
@app.post("/api/admin/students")
def create_student(body:StudentNew,who=Depends(admin_auth),db=Depends(getdb)):
    mat=body.matricula.strip()
    if db.scalar(select(Student).where(Student.matricula==mat)): raise HTTPException(409,"La matrícula ya existe.")
    s=Student(matricula=mat,nombre=body.nombre.strip(),password_hash=hashpw(body.password)); db.add(s); db.commit(); return {"ok":True,"matricula":mat}
@app.post("/api/admin/students/{sid}/toggle")
def toggle_student(sid:int,who=Depends(admin_auth),db=Depends(getdb)):
    s=db.get(Student,sid)
    if not s: raise HTTPException(404,"Alumno no encontrado.")
    s.active=not s.active; db.commit(); return {"ok":True,"active":s.active}
@app.post("/api/admin/qr-tokens")
def create_qr(body:QRNew,who=Depends(admin_auth),db=Depends(getdb)):
    s=db.scalar(select(Student).where(Student.matricula==body.matricula.strip(),Student.active.is_(True)))
    if not s: raise HTTPException(404,"Matrícula activa no encontrada.")
    raw=secrets.token_urlsafe(32); q=QRToken(token_hash=hashlib.sha256(raw.encode()).hexdigest(),student_id=s.id,expires_at=now()+timedelta(seconds=90)); db.add(q); db.commit()
    buf=io.BytesIO(); qrcode.make("SIGACEcyTEO:"+raw).save(buf,format="PNG")
    return {"matricula":s.matricula,"qr_data_url":"data:image/png;base64,"+base64.b64encode(buf.getvalue()).decode(),"expires_in_seconds":90}
@app.post("/api/admin/incidents/{iid}/resolve")
def resolve(iid:int,body:ResolveBody,who=Depends(admin_auth),db=Depends(getdb)):
    x=db.get(Incident,iid)
    if not x: raise HTTPException(404,"Incidencia no encontrada.")
    x.state="Resuelta"; x.resolved_at=now(); x.resolution_note=body.note[:1000]; db.commit(); return {"ok":True}
@app.post("/api/admin/commands")
def remote_command(body:CommandNew,who=Depends(admin_auth),db=Depends(getdb)):
    cmd=body.command.lower()
    if cmd not in ("logout","lock","unlock"): raise HTTPException(400,"Comando no válido.")
    c=db.scalar(select(Computer).where(Computer.code==body.computer_code.strip().upper()))
    if not c: raise HTTPException(404,"Equipo no registrado; inicia sesión en el cliente primero.")
    if cmd=="lock": c.state="Bloqueado"
    if cmd=="unlock": c.state="En uso" if db.scalar(select(LabSession).where(LabSession.computer_id==c.id,LabSession.ended_at.is_(None))) else "Libre"
    x=Command(computer_id=c.id,command=cmd,status="Pendiente"); db.add(x); db.commit(); return {"ok":True,"command_id":x.id}
@app.get("/api/admin/export/{kind}")
def export_report(kind:str,who=Depends(admin_auth),db=Depends(getdb)):
    if kind not in ("csv","xlsx","pdf"): raise HTTPException(400,"Formatos permitidos: csv, xlsx, pdf.")
    sessions=db.scalars(select(LabSession).order_by(LabSession.id)).all(); incidents=db.scalars(select(Incident).order_by(Incident.id)).all()
    cm={c.id:c.code for c in db.scalars(select(Computer)).all()}; sm={s.id:(s.matricula,s.nombre) for s in db.scalars(select(Student)).all()}
    heads=["Tipo","ID","Matrícula","Alumno","Equipo","Fecha inicio","Fecha fin","Estado","Detalle"]; rows=[]
    for s in sessions:
        mat,name=sm.get(s.student_id,("","")); rows.append(["Sesión",s.id,mat,name,cm.get(s.computer_id,""),ts(s.started_at),ts(s.ended_at),s.state,s.mode])
    for x in incidents:
        mat,name=sm.get(x.student_id,("","")); rows.append(["Incidencia",x.id,mat,name,cm.get(x.computer_id,""),ts(x.created_at),ts(x.resolved_at),x.state,x.category+": "+x.description])
    if kind=="csv":
        out=io.StringIO(); w=csv.writer(out); w.writerow(heads); w.writerows(rows); data=("\ufeff"+out.getvalue()).encode()
        return StreamingResponse(io.BytesIO(data),media_type="text/csv; charset=utf-8",headers={"Content-Disposition":'attachment; filename="sigacecyteo-bitacora.csv"'})
    if kind=="xlsx":
        wb=Workbook(); ws=wb.active; ws.title="Bitácora"; ws.append(heads)
        for row in rows: ws.append([str(v or "") for v in row])
        for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="12345B")
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(45,max(12,max(len(str(c.value or "")) for c in col)+2))
        out=io.BytesIO(); wb.save(out); out.seek(0)
        return StreamingResponse(out,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":'attachment; filename="sigacecyteo-bitacora.xlsx"'})
    out=io.BytesIO(); doc=SimpleDocTemplate(out,pagesize=landscape(letter),rightMargin=18,leftMargin=18,topMargin=20,bottomMargin=20); st=getSampleStyleSheet()
    data=[[Paragraph(str(v),st["BodyText"]) for v in heads]]
    for row in rows: data.append([Paragraph(str(v or "").replace("&","&amp;").replace("<","&lt;"),st["BodyText"]) for v in row])
    t=Table(data,repeatRows=1); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#12345B")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.3,colors.grey),("FONTSIZE",(0,0),(-1,-1),7),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.whitesmoke])]))
    doc.build([Paragraph("SIGACEcyTEO · Bitácora",st["Title"]),Spacer(1,8),t]); out.seek(0)
    return StreamingResponse(out,media_type="application/pdf",headers={"Content-Disposition":'attachment; filename="sigacecyteo-bitacora.pdf"'})
@app.get("/",response_class=HTMLResponse)
def panel(): return HTMLResponse(PANEL)

@app.websocket("/ws/client")
async def client_ws(ws:WebSocket,token:str):
    ident=TOKENS.get(token)
    if not ident or ident.get("kind")!="client" or ident.get("expires",0)<time.time(): await ws.close(code=1008); return
    await ws.accept()
    try:
        while True:
            ident=TOKENS.get(token)
            if not ident or ident.get("expires",0)<time.time(): break
            with SessionLocal() as db:
                live=db.get(LabSession,ident["session_id"])
                if not live or live.ended_at: break
                cmds=db.scalars(select(Command).where(Command.computer_id==ident["computer_id"],Command.status=="Pendiente").order_by(Command.id)).all()
                send=[]
                for x in cmds: x.status="Enviado"; x.sent_at=now(); send.append({"command_id":x.id,"command":x.command})
                db.commit()
            for x in send: await ws.send_json(x)
            try: await asyncio.wait_for(ws.receive_text(),timeout=1)
            except asyncio.TimeoutError: pass
    except (WebSocketDisconnect,RuntimeError): pass
    finally:
        try: await ws.close()
        except Exception: pass

PANEL=r"""<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>SIGACEcyTEO · Servidor</title><style>
*{box-sizing:border-box}body{margin:0;background:#f2f5fa;color:#172b45;font:15px Segoe UI,Arial}header{background:#12345b;color:white;padding:20px 5%;display:flex;justify-content:space-between;flex-wrap:wrap}main{max-width:1200px;margin:22px auto;padding:0 16px}.card{background:white;border:1px solid #e0e7f0;border-radius:12px;padding:16px;margin-bottom:14px}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:12px}.row{display:flex;gap:8px;flex-wrap:wrap;align-items:center}button,input{font:inherit;padding:10px;border:1px solid #cbd5e1;border-radius:7px}button{cursor:pointer;background:#e2eaf5;color:#12345b;font-weight:600}.primary{background:#175cd3;color:white}.danger{background:#b42318;color:white}.hidden{display:none!important}.muted{color:#66768b;font-size:13px}.num{font-size:28px;font-weight:800}table{width:100%;border-collapse:collapse;font-size:13px}td,th{padding:9px;border-bottom:1px solid #e3e9f2;text-align:left;vertical-align:top}th{background:#f6f8fc}.scroll{overflow:auto}label{font-size:13px;display:block;margin:10px 0 4px}#login{max-width:420px;margin:40px auto}#login input{width:100%}img.qr{width:190px;height:190px}.warn{padding:10px;background:#fff4d5;border-radius:8px;margin-bottom:12px}</style><body><header><div><b style="font-size:23px">SIGACEcyTEO</b><div class="muted" style="color:#c8d9ed">Servidor central · Acceso y bitácora</div></div><div id="who"></div></header><main>
<div id="login" class="card"><h2>Acceso de encargado</h2><label>Usuario</label><input id="user" autocomplete="username"><label>Contraseña</label><input id="pass" type="password" autocomplete="current-password"><p><button class="primary" onclick="doLogin()">Entrar al panel</button></p><div id="loginmsg" class="muted"></div><p class="muted">Administrador configurado en el .env del servidor.</p></div>
<div id="app" class="hidden"><div class="warn">"Bloquear" restringe la pantalla de la aplicación cliente; para bloqueo de Windows hay que desplegar el modo quiosco del sistema operativo.</div><div class="row"><button onclick="refresh()">Actualizar</button><button onclick="download('csv')">Exportar CSV</button><button onclick="download('xlsx')">Exportar Excel</button><button onclick="download('pdf')">Exportar PDF</button><button onclick="logoutAdmin()">Salir</button></div><div id="kpis" class="grid"></div>
<div class="card"><h2>Equipos</h2><div id="computers" class="grid"></div></div>
<div class="card"><h2>Incidencias</h2><div class="scroll"><table><thead><tr><th>Fecha</th><th>Equipo</th><th>Matrícula</th><th>Descripción</th><th>Estado</th><th></th></tr></thead><tbody id="incidents"></tbody></table></div></div>
<div class="card"><h2>Bitácora</h2><div class="scroll"><table><thead><tr><th>ID</th><th>Matrícula</th><th>Alumno</th><th>Equipo</th><th>Inicio</th><th>Fin</th><th>Modo</th><th>Estado</th></tr></thead><tbody id="sessions"></tbody></table></div></div>
<div class="grid"><div class="card"><h2>Registrar alumno</h2><label>Matrícula</label><input id="newmat"><label>Nombre</label><input id="newname"><label>Contraseña (mín. 10 caracteres)</label><input id="newpass" type="password"><p><button class="primary" onclick="createStudent()">Crear alumno</button></p><div id="studentmsg" class="muted"></div></div><div class="card"><h2>QR (un solo uso)</h2><label>Matrícula</label><input id="qrmat"><p><button class="primary" onclick="makeQR()">Generar QR (90 s)</button></p><div id="qrout"></div></div></div>
<div class="card"><h2>Alumnos</h2><div class="scroll"><table><thead><tr><th>Matrícula</th><th>Nombre</th><th>Estado</th><th></th></tr></thead><tbody id="students"></tbody></table></div></div><div class="card"><h2>Comandos recientes</h2><div id="commands" class="muted"></div></div></div></main>
<script>
let token='',username='';
function esc(v){return String(v??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]))}
async function api(p,m='GET',b=null){let r=await fetch(p,{method:m,headers:{'Content-Type':'application/json','Authorization':'Bearer '+token},body:b?JSON.stringify(b):undefined});let d=await r.json().catch(()=>({detail:'Respuesta no válida'}));if(!r.ok)throw Error(d.detail||'Error');return d}
async function doLogin(){try{let r=await fetch('/api/admin/login',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({username:document.getElementById('user').value,password:document.getElementById('pass').value})}),d=await r.json();if(!r.ok)throw Error(d.detail||'Acceso denegado');token=d.access_token;username=d.username;document.getElementById('login').classList.add('hidden');document.getElementById('app').classList.remove('hidden');document.getElementById('who').textContent='Encargado: '+username;refresh()}catch(e){document.getElementById('loginmsg').textContent=e.message}}
function logoutAdmin(){token='';document.getElementById('app').classList.add('hidden');document.getElementById('login').classList.remove('hidden')}
function fmt(v){return v?new Date(v).toLocaleString():'—'}
async function refresh(){if(!token)return;try{let d=await api('/api/admin/state');
document.getElementById('kpis').innerHTML=[['Equipos',d.computers.length],['En uso',d.computers.filter(x=>x.state==='En uso').length],['Sesiones activas',d.sessions.filter(x=>!x.ended_at).length],['Fallas abiertas',d.incidents.filter(x=>x.state==='Abierta').length]].map(x=>'<div class="card"><div class="muted">'+x[0]+'</div><div class="num">'+x[1]+'</div></div>').join('');
document.getElementById('computers').innerHTML=d.computers.map(c=>'<div class="card"><div class="row"><b>'+esc(c.code)+'</b> <b>'+esc(c.state)+'</b></div><p class="muted">Host: '+esc(c.hostname||'—')+' · IP: '+esc(c.ip_address||'—')+'<br>SO: '+esc(c.os_name||'—')+'<br>CPU '+esc(c.cpu??'—')+'% · RAM '+esc(c.memory??'—')+'% · Disco '+esc(c.disk??'—')+'%<br>Señal: '+esc(fmt(c.last_seen))+'</p><div class="row"><button onclick="sendCmd('+JSON.stringify(c.code)+',\'logout\')">Cerrar sesión</button><button onclick="sendCmd('+JSON.stringify(c.code)+',\'lock\')">Bloquear app</button><button onclick="sendCmd('+JSON.stringify(c.code)+',\'unlock\')">Desbloquear app</button></div></div>').join('')||'<p class="muted">Aún no hay equipos registrados.</p>';
document.getElementById('incidents').innerHTML=d.incidents.map(x=>'<tr><td>'+esc(fmt(x.created_at))+'</td><td>'+esc(x.computer_code)+'</td><td>'+esc(x.matricula)+'</td><td>'+esc(x.category+': '+x.description)+'</td><td>'+esc(x.state)+'</td><td>'+(x.state==='Abierta'?'<button onclick="resolveInc('+x.id+')">Resolver</button>':'—')+'</td></tr>').join('')||'<tr><td colspan="6">Sin incidencias</td></tr>';
document.getElementById('sessions').innerHTML=d.sessions.map(x=>'<tr><td>'+x.id+'</td><td>'+esc(x.matricula)+'</td><td>'+esc(x.nombre)+'</td><td>'+esc(x.computer_code)+'</td><td>'+esc(fmt(x.started_at))+'</td><td>'+esc(fmt(x.ended_at))+'</td><td>'+esc(x.mode)+'</td><td>'+esc(x.state)+'</td></tr>').join('');
document.getElementById('students').innerHTML=d.students.map(x=>'<tr><td>'+esc(x.matricula)+'</td><td>'+esc(x.nombre)+'</td><td>'+(x.active?'Activo':'Inactivo')+'</td><td><button onclick="toggleStudent('+x.id+')">'+(x.active?'Desactivar':'Activar')+'</button></td></tr>').join('');
document.getElementById('commands').innerHTML=d.commands.map(x=>esc(fmt(x.created_at)+' · '+x.computer_code+' · '+x.command+' · '+x.status)).join('<br>')||'Sin comandos'
}catch(e){console.error(e)}}
async function sendCmd(code,command){if(!confirm('¿Enviar '+command+' a '+code+'?'))return;try{await api('/api/admin/commands','POST',{computer_code:code,command});await refresh()}catch(e){alert(e.message)}}
async function resolveInc(id){let note=prompt('Nota de resolución (opcional):','Atendida por el encargado')||'';try{await api('/api/admin/incidents/'+id+'/resolve','POST',{note});await refresh()}catch(e){alert(e.message)}}
async function createStudent(){try{await api('/api/admin/students','POST',{matricula:document.getElementById('newmat').value,nombre:document.getElementById('newname').value,password:document.getElementById('newpass').value});document.getElementById('studentmsg').textContent='Alumno creado.';await refresh()}catch(e){document.getElementById('studentmsg').textContent=e.message}}
async function toggleStudent(id){try{await api('/api/admin/students/'+id+'/toggle','POST',{});await refresh()}catch(e){alert(e.message)}}
async function makeQR(){let o=document.getElementById('qrout');o.textContent='Generando…';try{let d=await api('/api/admin/qr-tokens','POST',{matricula:document.getElementById('qrmat').value});o.innerHTML='<p>QR para '+esc(d.matricula)+' (90 s)</p><img class="qr" alt="QR de acceso" src="'+d.qr_data_url+'">'}catch(e){o.textContent=e.message}}
async function download(k){try{let r=await fetch('/api/admin/export/'+k,{headers:{'Authorization':'Bearer '+token}});if(!r.ok)throw Error('Error al generar archivo');let blob=await r.blob(),url=URL.createObjectURL(blob),a=document.createElement('a');a.href=url;a.download='sigacecyteo-bitacora.'+k;a.click();URL.revokeObjectURL(url)}catch(e){alert(e.message)}}
setInterval(()=>{if(!document.getElementById('app').classList.contains('hidden'))refresh()},5000);
</script></body></html>"""
if __name__=="__main__": uvicorn.run(app,host=HOST,port=PORT,log_level="info")
